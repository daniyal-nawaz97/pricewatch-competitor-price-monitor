"""PriceWatch: know every time a competitor changes their price."""
from __future__ import annotations

import csv
import io
import logging
import re
import secrets
import threading
import time
from datetime import date, datetime, timedelta
from pathlib import Path

import httpx
from fastapi import Depends, FastAPI, File, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, PlainTextResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from . import alerts, analytics, collector, db, demo, matching, notify, reports, seed
from .config import (APP_NAME, BRAND_DIR, CONTACT_EMAIL, CONTACT_WHATSAPP, DEMO_EMAIL, MANUAL_SECONDS_PER_CHECK, SELF_URL, STATIC_DIR)

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
log = logging.getLogger("pricewatch")
app = FastAPI(title=APP_NAME, docs_url=None, redoc_url=None)
COOKIE = "pw_session"
STATE = {"ready": False, "status": "starting"}


# --------------------------------------------------------------------------- startup & scheduler
@app.on_event("startup")
def startup():
    db.init_schema()
    threading.Thread(target=_boot, daemon=True).start()


def _boot():
    try:
        from .config import ADMIN_EMAIL, ADMIN_PASSWORD, DEMO_MODE
        created = False
        if DEMO_MODE:
            STATE["status"] = "creating demo store"
            created = seed.seed_demo()
        elif ADMIN_EMAIL and ADMIN_PASSWORD and not db.one("SELECT id FROM users LIMIT 1"):
            db.execute("INSERT INTO users(email, name, role, password_hash, created_at) VALUES(?,?,?,?,?)",
                       (ADMIN_EMAIL, "Admin", "admin", db.hash_password(ADMIN_PASSWORD), db.now()))
        STATE["ready"] = True
        # wait for the web server to answer (the demo shops are served by this same app), then collect today live
        for _ in range(60):
            try:
                if httpx.get(SELF_URL + "/health", timeout=2).status_code == 200:
                    break
            except httpx.HTTPError:
                time.sleep(0.5)
        if (created or not db.one("SELECT id FROM runs WHERE day=? AND trigger!='demo-backfill'", (date.today().isoformat(),))) \
                and db.one("SELECT id FROM listings WHERE status='confirmed' LIMIT 1"):
            STATE["status"] = "collecting today's prices"
            collector.run(trigger="startup", deliver=False)
        STATE["status"] = "ready"
    except Exception:
        log.exception("startup failed")
        STATE["status"] = "error"
    _scheduler()


def _scheduler():
    """Every minute: run the daily check at check_time and send reports at report_time."""
    while True:
        try:
            s = db.get_settings()
            now = datetime.now()
            today = date.today().isoformat()
            if now.strftime("%H:%M") >= s["check_time"] and not db.one("SELECT id FROM runs WHERE day=? AND trigger IN ('schedule','startup','manual')", (today,)):
                collector.run(trigger="schedule")
            if now.strftime("%H:%M") >= s["report_time"]:
                for r in db.query("SELECT id FROM reports WHERE day=? AND sent_to IS NULL", (today,)):
                    reports.send(r["id"])
        except Exception:
            log.exception("scheduler error")
        time.sleep(60)


def today() -> str:
    return date.today().isoformat()


def data_day() -> str:
    """Latest day with collected prices (today once the morning check has run)."""
    r = db.one("SELECT MAX(day) AS d FROM observations WHERE ok=1")
    return r["d"] or today()


# --------------------------------------------------------------------------- auth
def current_user(request: Request):
    token = request.cookies.get(COOKIE)
    u = db.one("SELECT u.id, u.email, u.name, u.role FROM sessions s JOIN users u ON u.id=s.user_id WHERE s.token=?", (token,)) if token else None
    if not u:
        raise HTTPException(401, "Please sign in")
    return u


def need(*roles):
    def dep(user=Depends(current_user)):
        if user["role"] not in roles:
            raise HTTPException(403, "You don't have permission for this. Ask an admin.")
        return user
    return dep


EDIT = need("admin", "analyst")


def _login(response, uid):
    token = secrets.token_urlsafe(32)
    db.execute("INSERT INTO sessions(token, user_id, created_at) VALUES(?,?,?)", (token, uid, db.now()))
    response.set_cookie(COOKIE, token, httponly=True, samesite="lax", max_age=60 * 60 * 24 * 14)


class LoginIn(BaseModel):
    email: str
    password: str


@app.post("/api/login")
def login(body: LoginIn, response: Response):
    u = db.one("SELECT * FROM users WHERE lower(email)=lower(?)", (body.email.strip(),))
    if not u or not db.check_password(body.password, u["password_hash"]):
        raise HTTPException(401, "Email or password is not correct")
    _login(response, u["id"])
    return {"ok": True}


@app.post("/api/demo-login")
def demo_login(response: Response):
    u = db.one("SELECT id FROM users WHERE email=?", (DEMO_EMAIL,))
    if not u or not STATE["ready"]:
        raise HTTPException(503, "The demo store is still being prepared. Try again in a few seconds.")
    _login(response, u["id"])
    return {"ok": True}


@app.post("/api/logout")
def logout(request: Request, response: Response):
    db.execute("DELETE FROM sessions WHERE token=?", (request.cookies.get(COOKIE, ""),))
    response.delete_cookie(COOKIE)
    return {"ok": True}


@app.get("/api/me")
def me(user=Depends(current_user)):
    return user


@app.get("/api/app-info")
def app_info():
    s = db.get_settings() if db.one("SELECT name FROM sqlite_master WHERE name='settings'") else db.DEFAULT_SETTINGS
    return {"app_name": APP_NAME, "company_name": s["company_name"], "accent_color": s["accent_color"], "logo_url": s["logo_url"],
            "currency": s["currency"], "rounding": s["rounding"], "ready": STATE["ready"], "status": STATE["status"],
            "contact_email": CONTACT_EMAIL, "contact_whatsapp": CONTACT_WHATSAPP}


# --------------------------------------------------------------------------- dashboard
@app.get("/api/overview")
def overview(user=Depends(current_user)):
    d = data_day()
    o = analytics.overview(d)
    run = collector.last_run()
    o["last_run"] = run
    o["unread_alerts"] = db.one("SELECT COUNT(*) AS n FROM alerts WHERE is_read=0")["n"]
    o["pending_matches"] = db.one("SELECT COUNT(*) AS n FROM listings WHERE status='pending'")["n"]
    o["is_demo"] = bool(db.one("SELECT id FROM competitors WHERE is_demo=1 LIMIT 1"))
    o["hours_saved_week"] = round(o["cards"]["listings"] * 7 * MANUAL_SECONDS_PER_CHECK / 3600, 1)
    o["manual_seconds"] = MANUAL_SECONDS_PER_CHECK
    return o


@app.get("/api/products")
def products(q: str = "", category: str = "", brand: str = "", only: str = "", sort: str = "name", user=Depends(current_user)):
    d = data_day()
    rows = analytics.snapshot(d)
    if q:
        ql = q.lower()
        rows = [r for r in rows if ql in r["name"].lower() or ql in r["sku"].lower()]
    if category:
        rows = [r for r in rows if r["category"] == category]
    if brand:
        rows = [r for r in rows if r["brand"] == brand]
    if only == "more_expensive":
        rows = [r for r in rows if r["gap"] is not None and r["gap"] > 0]
    elif only == "cheapest":
        rows = [r for r in rows if r["position"] == "cheapest"]
    elif only == "top":
        rows = [r for r in rows if r["top"]]
    if sort == "gap":
        rows.sort(key=lambda r: -(r["gap_pct"] if r["gap_pct"] is not None else -999))
    elif sort == "price":
        rows.sort(key=lambda r: -r["price"])
    comps = db.query("SELECT id, name, color, is_demo FROM competitors WHERE active=1 ORDER BY id")
    all_rows = db.query("SELECT DISTINCT category, brand FROM products WHERE active=1")
    return {"day": d, "competitors": comps, "rows": rows, "categories": sorted({r["category"] for r in all_rows}),
            "brands": sorted({r["brand"] for r in all_rows if r["brand"]})}


@app.get("/api/products/export.xlsx")
def export_products(user=Depends(current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    d = data_day()
    rows = analytics.snapshot(d)
    comps = db.query("SELECT id, name FROM competitors WHERE active=1 ORDER BY id")
    wb = Workbook()
    ws = wb.active
    ws.title = f"Prices {d}"
    ws.append(["SKU", "Product", "Category", "Your price"] + [c["name"] for c in comps] + ["Cheapest competitor", "Gap", "Gap %", "Position"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="4F46E5")
    for r in rows:
        cells = []
        for c in comps:
            x = r["competitors"].get(c["id"], {})
            cells.append(f"{x['price']:.0f}{'' if x.get('in_stock') else ' (out of stock)'}" if x.get("price") else "")
        ws.append([r["sku"], r["name"], r["category"], r["price"]] + cells + [r["cheapest"], r["gap"], round(r["gap_pct"], 1) if r["gap_pct"] is not None else None,
                                                                             r["position"].replace("_", " ")])
    for col, w in zip("ABCDEFGHIJKL", [10, 44, 16, 11, 16, 16, 16, 18, 10, 9, 16, 12]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="PriceWatch_prices_{d}.xlsx"'})


@app.get("/api/products/{pid}")
def product(pid: int, days: int = 30, user=Depends(current_user)):
    if days not in (7, 30, 90):
        days = 30
    p = analytics.product_detail(pid, data_day(), days)
    if not p:
        raise HTTPException(404, "Product not found")
    p["competitors"] = db.query("SELECT id, name, color FROM competitors WHERE active=1 ORDER BY id")
    return p


class ProductIn(BaseModel):
    sku: str
    name: str
    brand: str = ""
    category: str = ""
    price: float
    cost: float | None = None
    top: bool = False


@app.post("/api/products")
def add_product(body: ProductIn, user=Depends(EDIT)):
    if db.one("SELECT id FROM products WHERE sku=?", (body.sku.strip(),)):
        raise HTTPException(400, "A product with this SKU already exists")
    pid = db.execute("INSERT INTO products(sku, name, brand, category, price, cost, top, created_at) VALUES(?,?,?,?,?,?,?,?)",
                     (body.sku.strip(), body.name.strip(), body.brand.strip(), body.category.strip() or "Other", body.price, body.cost, int(body.top), db.now()))
    db.execute("INSERT OR REPLACE INTO own_prices(product_id, day, price) VALUES(?,?,?)", (pid, today(), body.price))
    return {"id": pid}


@app.put("/api/products/{pid}")
def update_product(pid: int, body: ProductIn, user=Depends(EDIT)):
    db.execute("UPDATE products SET sku=?, name=?, brand=?, category=?, price=?, cost=?, top=? WHERE id=?",
               (body.sku.strip(), body.name.strip(), body.brand.strip(), body.category.strip() or "Other", body.price, body.cost, int(body.top), pid))
    db.execute("INSERT OR REPLACE INTO own_prices(product_id, day, price) VALUES(?,?,?)", (pid, today(), body.price))
    return {"ok": True}


@app.delete("/api/products/{pid}")
def delete_product(pid: int, user=Depends(EDIT)):
    db.execute("UPDATE products SET active=0 WHERE id=?", (pid,))
    return {"ok": True}


@app.post("/api/products/import")
async def import_products(file: UploadFile = File(...), user=Depends(EDIT)):
    """CSV columns: sku, name, price (required) and brand, category, cost, top (optional)."""
    text = (await file.read()).decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    cols = {c.lower().strip(): c for c in (reader.fieldnames or [])}
    if not {"sku", "name", "price"} <= set(cols):
        raise HTTPException(400, "The CSV needs at least these columns: sku, name, price")
    added = updated = 0
    for row in reader:
        g = lambda k: (row.get(cols[k]) or "").strip() if k in cols else ""  # noqa: E731
        try:
            price = float(re.sub(r"[^\d.]", "", g("price")))
        except ValueError:
            continue
        cost = float(re.sub(r"[^\d.]", "", g("cost"))) if g("cost") else None
        ex = db.one("SELECT id FROM products WHERE sku=?", (g("sku"),))
        if ex:
            db.execute("UPDATE products SET name=?, brand=?, category=?, price=?, cost=?, active=1 WHERE id=?",
                       (g("name"), g("brand"), g("category") or "Other", price, cost, ex["id"]))
            pid, updated = ex["id"], updated + 1
        else:
            pid = db.execute("INSERT INTO products(sku, name, brand, category, price, cost, top, created_at) VALUES(?,?,?,?,?,?,?,?)",
                             (g("sku"), g("name"), g("brand"), g("category") or "Other", price, cost, int(g("top").lower() in ("1", "yes", "true")), db.now()))
            added += 1
        db.execute("INSERT OR REPLACE INTO own_prices(product_id, day, price) VALUES(?,?,?)", (pid, today(), price))
    return {"added": added, "updated": updated}


# --------------------------------------------------------------------------- alerts
@app.get("/api/alerts")
def list_alerts(type: str = "", competitor: int = 0, product: int = 0, limit: int = 100, user=Depends(current_user)):
    sql = ("SELECT a.*, p.name AS product, c.name AS competitor, c.color FROM alerts a LEFT JOIN products p ON p.id=a.product_id "
           "LEFT JOIN competitors c ON c.id=a.competitor_id WHERE 1=1")
    args = []
    if type:
        sql += " AND a.type=?"; args.append(type)
    if competitor:
        sql += " AND a.competitor_id=?"; args.append(competitor)
    if product:
        sql += " AND a.product_id=?"; args.append(product)
    sql += " ORDER BY a.day DESC, CASE a.severity WHEN 'high' THEN 0 WHEN 'warn' THEN 1 WHEN 'good' THEN 2 ELSE 3 END, a.id DESC LIMIT ?"
    args.append(min(limit, 500))
    return {"alerts": db.query(sql, args), "types": alerts.TYPES, "competitors": db.query("SELECT id, name FROM competitors ORDER BY id")}


@app.post("/api/alerts/read")
def read_alerts(user=Depends(current_user)):
    db.execute("UPDATE alerts SET is_read=1 WHERE is_read=0")
    return {"ok": True}


# --------------------------------------------------------------------------- competitors, listings, matches
@app.get("/api/competitors")
def competitors(user=Depends(current_user)):
    return {"competitors": analytics.competitor_health(data_day()), "last_run": collector.last_run()}


class CompetitorIn(BaseModel):
    name: str
    website: str = ""
    color: str = "#64748b"
    price_selector: str = ""
    stock_selector: str = ""
    catalog_url: str = ""
    active: bool = True


@app.post("/api/competitors")
def add_competitor(body: CompetitorIn, user=Depends(EDIT)):
    cid = db.execute("INSERT INTO competitors(name, website, color, price_selector, stock_selector, catalog_url, created_at) VALUES(?,?,?,?,?,?,?)",
                     (body.name.strip(), body.website.strip(), body.color, body.price_selector.strip() or None, body.stock_selector.strip() or None,
                      body.catalog_url.strip() or None, db.now()))
    return {"id": cid}


@app.put("/api/competitors/{cid}")
def update_competitor(cid: int, body: CompetitorIn, user=Depends(EDIT)):
    db.execute("UPDATE competitors SET name=?, website=?, color=?, price_selector=?, stock_selector=?, catalog_url=?, active=? WHERE id=?",
               (body.name.strip(), body.website.strip(), body.color, body.price_selector.strip() or None, body.stock_selector.strip() or None,
                body.catalog_url.strip() or None, int(body.active), cid))
    return {"ok": True}


class ListingIn(BaseModel):
    product_id: int
    competitor_id: int
    url: str


@app.post("/api/listings")
def add_listing(body: ListingIn, user=Depends(EDIT)):
    if not re.match(r"^(https?://|/demo-shops/)", body.url.strip()):
        raise HTTPException(400, "Please paste the full product page address, starting with https://")
    try:
        lid = db.execute("INSERT INTO listings(product_id, competitor_id, url, status, confidence, created_at) VALUES(?,?,?,?,?,?)",
                         (body.product_id, body.competitor_id, body.url.strip(), "confirmed", 1.0, db.now()))
    except Exception:
        raise HTTPException(400, "This page is already tracked")
    return test_listing(lid, user)


@app.post("/api/listings/{lid}/test")
def test_listing(lid: int, user=Depends(EDIT)):
    l = db.one("SELECT * FROM listings WHERE id=?", (lid,))
    c = db.one("SELECT * FROM competitors WHERE id=?", (l["competitor_id"],))
    with httpx.Client() as client:
        res = collector.check_listing(l, c, client)
    collector.record(lid, today(), res)
    if res["ok"] and res.get("title") and not l["title"]:
        db.execute("UPDATE listings SET title=? WHERE id=?", (res["title"], lid))
    return {"listing_id": lid, **res}


@app.delete("/api/listings/{lid}")
def delete_listing(lid: int, user=Depends(EDIT)):
    db.execute("UPDATE listings SET status='removed' WHERE id=?", (lid,))
    return {"ok": True}


@app.get("/api/matches")
def matches(user=Depends(current_user)):
    rows = db.query("""SELECT l.*, c.name AS competitor, c.color, p.name AS product, p.sku, p.price AS your_price, p.category
                       FROM listings l JOIN competitors c ON c.id=l.competitor_id LEFT JOIN products p ON p.id=l.product_id
                       WHERE l.status='pending' ORDER BY l.confidence DESC""")
    for r in rows:
        r["label"] = matching.label(r["confidence"] or 0) if r["product_id"] else "No match"
        o = db.one("SELECT price, in_stock FROM observations WHERE listing_id=? AND ok=1 ORDER BY day DESC LIMIT 1", (r["id"],))
        r["their_price"] = o["price"] if o else None
        if r["their_price"] is None:
            try:
                with httpx.Client() as client:
                    res = collector.check_listing(r, db.one("SELECT * FROM competitors WHERE id=?", (r["competitor_id"],)), client)
                if res["ok"]:
                    collector.record(r["id"], data_day(), res)
                    r["their_price"] = res["price"]
            except Exception:
                pass
    reviewed = db.one("SELECT COUNT(*) AS n FROM listings WHERE status IN ('confirmed','rejected')")["n"]
    return {"pending": rows, "reviewed": reviewed, "products": db.query("SELECT id, name, sku FROM products WHERE active=1 ORDER BY name")}


class MatchIn(BaseModel):
    decision: str  # yes | no | skip
    product_id: int | None = None


@app.post("/api/matches/{lid}")
def decide(lid: int, body: MatchIn, user=Depends(EDIT)):
    l = db.one("SELECT * FROM listings WHERE id=? AND status='pending'", (lid,))
    if not l:
        raise HTTPException(404, "Already reviewed")
    if body.decision == "yes":
        pid = body.product_id or l["product_id"]
        if not pid:
            raise HTTPException(400, "Choose which of your products this is")
        db.execute("UPDATE listings SET status='confirmed', product_id=? WHERE id=?", (pid, lid))
    elif body.decision == "no":
        db.execute("UPDATE listings SET status='rejected' WHERE id=?", (lid,))
    elif body.decision == "skip":
        db.execute("UPDATE listings SET confidence=COALESCE(confidence,0)-0.001 WHERE id=?", (lid,))
    else:
        raise HTTPException(400, "Decision must be yes, no or skip")
    return {"ok": True}


@app.post("/api/runs")
def run_now(user=Depends(EDIT)):
    r = collector.last_run()
    if r and r["status"] == "running":
        return {"status": "busy"}
    collector.run_async("manual")
    return {"status": "started"}


@app.get("/api/runs/latest")
def run_latest(user=Depends(current_user)):
    return collector.last_run()


# --------------------------------------------------------------------------- reports
@app.get("/api/reports")
def list_reports(user=Depends(current_user)):
    import json
    rows = db.query("SELECT id, kind, day, title, summary_json, sent_to, created_at FROM reports ORDER BY day DESC, kind LIMIT 60")
    for r in rows:
        r["summary"] = json.loads(r.pop("summary_json") or "{}")
    return {"reports": rows, "channels": notify.channels_status()}


@app.get("/api/reports/{rid}/{fmt}")
def report_file(rid: int, fmt: str, user=Depends(current_user)):
    r = db.one("SELECT * FROM reports WHERE id=?", (rid,))
    if not r or fmt not in ("html", "pdf", "xlsx"):
        raise HTTPException(404, "Report not found")
    path = Path(r[f"{fmt}_path"])
    if fmt == "html":
        return HTMLResponse(path.read_text())
    return FileResponse(path, filename=f"PriceWatch_{r['kind']}_{r['day']}.{fmt}")


@app.post("/api/reports/{rid}/send")
def send_report(rid: int, user=Depends(EDIT)):
    return {"sent_to": reports.send(rid)}


# --------------------------------------------------------------------------- settings & users
@app.get("/api/settings")
def get_settings(user=Depends(current_user)):
    s = db.get_settings()
    s["channel_status"] = notify.channels_status()
    s["categories"] = sorted({r["category"] for r in db.query("SELECT category FROM products WHERE active=1")})
    return s


@app.put("/api/settings")
def put_settings(body: dict, user=Depends(need("admin"))):
    if "accent_color" in body and not re.fullmatch(r"#[0-9a-fA-F]{6}", str(body["accent_color"])):
        raise HTTPException(400, "Colour must look like #4f46e5")
    for k in ("check_time", "report_time"):
        if k in body and not re.fullmatch(r"[0-2]\d:[0-5]\d", str(body[k])):
            raise HTTPException(400, "Times must look like 07:00")
    for k in ("daily_recipients", "weekly_recipients"):
        if k in body:
            body[k] = [e.strip() for e in body[k] if re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", e.strip())]
    db.set_settings({k: v for k, v in body.items() if k != "demo_anchor"})
    return get_settings(user)


@app.post("/api/settings/logo")
async def upload_logo(file: UploadFile = File(...), user=Depends(need("admin"))):
    ext = Path(file.filename or "").suffix.lower()
    if ext not in (".png", ".jpg", ".jpeg", ".svg", ".webp"):
        raise HTTPException(400, "Please upload a PNG, JPG, SVG or WEBP logo")
    name = f"logo_{secrets.token_hex(4)}{ext}"
    (BRAND_DIR / name).write_bytes(await file.read())
    db.set_settings({"logo_url": f"/brand/{name}"})
    return {"ok": True}


@app.get("/api/users")
def users(user=Depends(need("admin"))):
    return db.query("SELECT id, name, email, role FROM users ORDER BY id")


class UserIn(BaseModel):
    name: str
    email: str
    role: str = "viewer"


@app.post("/api/users")
def add_user(body: UserIn, user=Depends(need("admin"))):
    if body.role not in ("admin", "analyst", "viewer"):
        raise HTTPException(400, "Role must be admin, analyst or viewer")
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", body.email.strip()):
        raise HTTPException(400, "Please enter a valid email")
    if db.one("SELECT id FROM users WHERE lower(email)=lower(?)", (body.email.strip(),)):
        raise HTTPException(400, "This person already has access")
    temp = secrets.token_urlsafe(6)
    db.execute("INSERT INTO users(email, name, role, password_hash, created_at) VALUES(?,?,?,?,?)",
               (body.email.strip(), body.name.strip() or body.email, body.role, db.hash_password(temp), db.now()))
    return {"temporary_password": temp}


@app.patch("/api/users/{uid}")
def change_role(uid: int, body: dict, user=Depends(need("admin"))):
    if body.get("role") not in ("admin", "analyst", "viewer") or uid == user["id"]:
        raise HTTPException(400, "Can't change this role")
    db.execute("UPDATE users SET role=? WHERE id=?", (body["role"], uid))
    return {"ok": True}


@app.delete("/api/users/{uid}")
def remove_user(uid: int, user=Depends(need("admin"))):
    if uid == user["id"]:
        raise HTTPException(400, "You can't remove yourself")
    db.execute("DELETE FROM sessions WHERE user_id=?", (uid,))
    db.execute("DELETE FROM users WHERE id=?", (uid,))
    return {"ok": True}


# --------------------------------------------------------------------------- fictitious demo competitor shops
def _anchor() -> date:
    a = db.get_settings().get("demo_anchor")
    return date.fromisoformat(a) if a else date.today()


@app.get("/demo-shops/robots.txt")
def shops_robots():
    return PlainTextResponse("User-agent: *\nAllow: /\n")


@app.get("/demo-shops/{shop}/")
def shop_home(shop: str):
    if shop not in {c["slug"] for c in demo.COMPETITORS}:
        raise HTTPException(404)
    return HTMLResponse(demo.shop_catalog_html(shop, date.today(), _anchor()))


@app.get("/demo-shops/{shop}/p/{handle}")
def shop_product(shop: str, handle: str):
    html = demo.shop_product_html(shop, handle, date.today(), _anchor()) if shop in {c["slug"] for c in demo.COMPETITORS} else None
    if not html:
        raise HTTPException(404, "Product not found")
    return HTMLResponse(html)


# --------------------------------------------------------------------------- pages
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
app.mount("/brand", StaticFiles(directory=BRAND_DIR), name="brand")


@app.get("/health")
def health():
    return {"ok": True, "ready": STATE["ready"], "status": STATE["status"]}


@app.get("/landing")
def landing():
    return FileResponse(STATIC_DIR / "landing.html")


@app.get("/privacy")
def privacy():
    return FileResponse(STATIC_DIR / "privacy.html")


@app.get("/")
def index():
    return FileResponse(STATIC_DIR / "index.html", headers={"Cache-Control": "no-cache"})


@app.exception_handler(404)
async def not_found(request: Request, exc):
    if request.url.path.startswith(("/api/", "/demo-shops/")):
        return JSONResponse({"detail": getattr(exc, "detail", "Not found")}, status_code=404)
    return RedirectResponse("/")
